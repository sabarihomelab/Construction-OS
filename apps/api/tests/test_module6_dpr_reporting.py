import io

import pytest
from docx import Document
from fastapi import FastAPI

from app.db import model_registry as _model_registry  # noqa: F401
from app.db.base import Base
from app.modules.field.api import router
from app.modules.field.dpr_report_type import DPR_REPORT_TYPE_KEY
from app.modules.field.dpr_schemas import DPRRenderRequest
from app.modules.files.object_store import LocalGeneratedObjectStore
from app.modules.reporting.template_execution import render_uploaded_docx
from app.modules.reporting.template_models import TemplateOutputFormat
from app.modules.reporting.template_renderer import render_report
from app.modules.reporting.template_schemas import (
    ReportBlock,
    ReportBlockKind,
    ReportLayoutSpec,
    TableColumnSpec,
)
from app.modules.reporting.type_registry import (
    AttachmentPolicy,
    LocationStampPolicy,
    ReportGenerationTrigger,
    report_types,
)


def _paths() -> dict[str, dict[str, object]]:
    application = FastAPI()
    application.include_router(router, prefix="/api/v1")
    return application.openapi()["paths"]


def _layout() -> ReportLayoutSpec:
    return ReportLayoutSpec(
        report_title="{{project.name}} - DPR",
        blocks=[
            ReportBlock(
                id="workforce",
                kind=ReportBlockKind.TABLE,
                order=10,
                row=0,
                width_percent=100,
                title="Workforce",
                data_key="workforce",
                table_columns=[
                    TableColumnSpec(key="trade", label="Trade", order=10),
                    TableColumnSpec(key="worker_count", label="Workers", order=20),
                ],
            )
        ],
    )


def test_dpr_report_type_declares_business_specific_capabilities() -> None:
    contract = report_types.get(DPR_REPORT_TYPE_KEY)

    assert contract.default_trigger == ReportGenerationTrigger.ON_APPROVAL
    assert ReportGenerationTrigger.MANUAL in contract.allowed_triggers
    assert contract.attachment_policy == AttachmentPolicy.OPTIONAL
    assert contract.location_stamp_policy == LocationStampPolicy.OPTIONAL
    assert contract.supports_photos is True
    assert contract.supports_signatures is True
    assert contract.issued_output_formats == ("pdf", "docx")
    assert "{{project.number}}" in contract.filename_pattern


def test_preview_is_default_and_official_issue_is_explicit() -> None:
    request = DPRRenderRequest(output_format=TemplateOutputFormat.PDF)
    assert request.persist_history is False
    assert request.generation_trigger == ReportGenerationTrigger.MANUAL


def test_dpr_api_exposes_work_progress_issue_history_and_download() -> None:
    paths = _paths()
    root = "/api/v1/projects/{project_id}/daily-reports/{report_id}"

    assert set(paths[f"{root}/work-progress"]) >= {"get", "put"}
    assert f"{root}/report-payload" in paths
    assert f"{root}/render" in paths
    assert f"{root}/render-history" in paths
    assert f"{root}/render-history/{{render_id}}/download" in paths
    assert "/api/v1/projects/{project_id}/dpr-templates" in paths


def test_report_history_pins_exact_file_version() -> None:
    table = Base.metadata.tables["report_render_records"]
    assert "output_filename" in table.c
    assert "output_file_asset_id" in table.c
    assert "output_file_version" in table.c
    assert "generation_trigger" in table.c
    assert "issued_at" in table.c

    targets = {
        fk.target_fullname
        for fk in table.c.output_file_version.foreign_keys
    }
    assert "file_versions.version" in targets


def test_uploaded_word_template_injects_scalars_and_collection_table() -> None:
    source = Document()
    source.add_paragraph("Project: {{project.name}}")
    source.add_paragraph("{{workforce}}")
    buffer = io.BytesIO()
    source.save(buffer)

    payload = {
        "project": {"name": "Metro Package A"},
        "workforce": [
            {"trade": "Bar Bender", "worker_count": 12},
            {"trade": "Carpenter", "worker_count": 8},
        ],
    }
    rendered = render_uploaded_docx(buffer.getvalue(), payload, _layout())
    result = Document(io.BytesIO(rendered))

    assert any("Metro Package A" in paragraph.text for paragraph in result.paragraphs)
    assert len(result.tables) == 1
    assert result.tables[0].cell(0, 0).text == "Trade"
    assert result.tables[0].cell(1, 0).text == "Bar Bender"
    assert result.tables[0].cell(2, 1).text == "8"


def test_pdf_renderer_handles_long_tables_with_one_normalized_payload() -> None:
    payload = {
        "project": {"name": "Long Table Project"},
        "workforce": [
            {"trade": f"Trade {index}", "worker_count": index}
            for index in range(1, 180)
        ],
    }
    content, content_type, extension = render_report(
        payload,
        _layout(),
        TemplateOutputFormat.PDF,
    )

    assert content.startswith(b"%PDF")
    assert content_type == "application/pdf"
    assert extension == "pdf"


@pytest.mark.asyncio
async def test_local_generated_object_store_persists_exact_bytes(tmp_path) -> None:
    store = LocalGeneratedObjectStore(tmp_path)
    stored = await store.write_bytes(
        organization_id="00000000-0000-0000-0000-000000000001",
        filename="DPR-001.pdf",
        content=b"%PDF-test",
        content_type="application/pdf",
    )

    assert stored.provider_key == "local"
    assert await store.read_bytes(stored.storage_key) == b"%PDF-test"
