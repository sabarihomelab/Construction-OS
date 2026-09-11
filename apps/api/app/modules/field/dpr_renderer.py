import csv
import io
import json
from collections.abc import Mapping, Sequence
from html import escape
from typing import Any

from docx import Document
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, LETTER, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.modules.field.dpr_reporting_models import DPROutputFormat
from app.modules.field.dpr_reporting_schemas import DPRLayoutElementKind, DPRLayoutSpec


_CONTENT_TYPES = {
    DPROutputFormat.PDF: "application/pdf",
    DPROutputFormat.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    DPROutputFormat.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    DPROutputFormat.CSV: "text/csv; charset=utf-8",
    DPROutputFormat.HTML: "text/html; charset=utf-8",
    DPROutputFormat.JSON: "application/json; charset=utf-8",
}


_EXTENSIONS = {item: item.value for item in DPROutputFormat}


_DEFAULT_FIELDS: dict[str, tuple[str, ...]] = {
    "crew": ("company_name", "trade", "worker_count", "regular_hours", "overtime_hours", "notes"),
    "work": ("description", "location", "cost_code", "quantity", "unit_code", "notes"),
    "materials": ("material_name", "quantity", "unit_code", "location", "notes"),
    "equipment": ("equipment_name", "equipment_reference", "hours_operated", "status", "notes"),
    "deliveries": ("supplier", "material", "quantity", "unit_code", "delivered_at", "ticket_number", "notes"),
    "production": ("description", "cost_code", "quantity", "unit_code", "location", "notes"),
    "delays": ("category", "description", "started_at", "ended_at", "lost_hours", "responsible_party", "schedule_impact", "notes"),
    "safety": ("entry_type", "summary", "severity", "notes"),
    "photos": ("asset_id", "pinned_version", "relation_type"),
}


def _company_name(snapshot: Mapping[str, Any], layout: DPRLayoutSpec) -> str:
    organization = snapshot.get("organization", {})
    if not isinstance(organization, Mapping):
        organization = {}
    mode = layout.branding.company_name_mode.value
    if mode == "custom":
        return (layout.branding.custom_company_name or "").strip()
    if mode == "legal":
        return str(organization.get("legal_name") or organization.get("name") or "")
    return str(organization.get("name") or "")


def _ordered_elements(layout: DPRLayoutSpec):
    return sorted((item for item in layout.elements if item.visible), key=lambda item: (item.order, item.id))


def _rows(snapshot: Mapping[str, Any], key: str) -> list[Mapping[str, Any]]:
    sections = snapshot.get("sections", {})
    if not isinstance(sections, Mapping):
        return []
    value = sections.get(key, [])
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes)):
        return []
    return [row for row in value if isinstance(row, Mapping)]


def _field_list(element, section_key: str, rows: list[Mapping[str, Any]]) -> list[str]:
    if element.fields:
        return list(element.fields)
    configured = _DEFAULT_FIELDS.get(section_key)
    if configured:
        return list(configured)
    if not rows:
        return []
    return [key for key in rows[0] if not key.endswith("_id") and key not in {"id", "source"}]


def _label(key: str) -> str:
    return key.replace("_", " ").strip().title()


def _scalar(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _project_summary(snapshot: Mapping[str, Any]) -> list[tuple[str, str]]:
    project = snapshot.get("project", {})
    report = snapshot.get("report", {})
    if not isinstance(project, Mapping):
        project = {}
    if not isinstance(report, Mapping):
        report = {}
    return [
        ("Project", str(project.get("name") or "")),
        ("Project Number", str(project.get("number") or "")),
        ("Report Date", str(report.get("report_date") or "")),
        ("Shift", str(report.get("shift_code") or "")),
        ("Status", str(report.get("status") or "")),
        ("Revision", str(report.get("revision") or "")),
    ]


def _weather(snapshot: Mapping[str, Any]) -> list[tuple[str, str]]:
    report = snapshot.get("report", {})
    if not isinstance(report, Mapping):
        report = {}
    unit = str(report.get("temperature_unit") or "")
    low = report.get("temperature_low")
    high = report.get("temperature_high")
    return [
        ("Condition", _scalar(report.get("weather_condition"))),
        ("Low", f"{_scalar(low)} {unit}".strip()),
        ("High", f"{_scalar(high)} {unit}".strip()),
    ]


def _html_table(rows: list[Mapping[str, Any]], fields: list[str]) -> str:
    if not rows:
        return '<p class="empty">No entries</p>'
    heads = "".join(f"<th>{escape(_label(field))}</th>" for field in fields)
    body = []
    for row in rows:
        cells = "".join(f"<td>{escape(_scalar(row.get(field)))}</td>" for field in fields)
        body.append(f"<tr>{cells}</tr>")
    return f"<table><thead><tr>{heads}</tr></thead><tbody>{''.join(body)}</tbody></table>"


def _html(snapshot: Mapping[str, Any], layout: DPRLayoutSpec) -> bytes:
    page = layout.page
    company = escape(_company_name(snapshot, layout))
    project = snapshot.get("project", {}) if isinstance(snapshot.get("project"), Mapping) else {}
    report = snapshot.get("report", {}) if isinstance(snapshot.get("report"), Mapping) else {}
    title = escape(layout.branding.report_title)
    logo_marker = ""
    if layout.branding.logo_asset_id:
        logo_marker = (
            f'<div class="logo" data-asset-id="{layout.branding.logo_asset_id}" '
            f'data-version="{layout.branding.logo_version}">Company logo</div>'
        )

    blocks: list[str] = []
    for element in _ordered_elements(layout):
        if element.kind == DPRLayoutElementKind.PAGE_BREAK:
            blocks.append('<div class="page-break"></div>')
            continue
        if element.start_new_page:
            blocks.append('<div class="page-break"></div>')
        title_text = escape(element.title or _label(element.kind.value))
        content = ""
        if element.kind == DPRLayoutElementKind.PROJECT_SUMMARY:
            content = _html_table(
                [{"field": key, "value": value} for key, value in _project_summary(snapshot)],
                ["field", "value"],
            )
        elif element.kind == DPRLayoutElementKind.WEATHER:
            content = _html_table(
                [{"field": key, "value": value} for key, value in _weather(snapshot)],
                ["field", "value"],
            )
        elif element.kind == DPRLayoutElementKind.NOTES:
            content = f"<p>{escape(_scalar(report.get('notes')))}</p>"
        elif element.kind == DPRLayoutElementKind.CUSTOM_TEXT:
            content = f"<p>{escape(element.custom_text or '')}</p>"
        elif element.kind == DPRLayoutElementKind.SIGNATURES:
            content = '<div class="signatures"><span>Prepared by</span><span>Checked by</span><span>Approved by</span></div>'
        else:
            section_key = element.kind.value
            rows = _rows(snapshot, section_key)
            content = _html_table(rows, _field_list(element, section_key, rows))
        blocks.append(
            f'<section class="span-{element.column_span}"><h2>{title_text}</h2>{content}</section>'
        )

    orientation = page.orientation.value
    css_page = f"{page.size.value.upper()} {orientation}"
    header = layout.header_footer
    header_html = " | ".join(filter(None, [header.header_left, header.header_center, header.header_right]))
    footer_html = " | ".join(filter(None, [header.footer_left, header.footer_center, header.footer_right]))
    html = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{title}</title>
<style>
@page {{ size: {css_page}; margin: {page.margin_top_mm}mm {page.margin_right_mm}mm {page.margin_bottom_mm}mm {page.margin_left_mm}mm; }}
body {{ font-family: Arial, sans-serif; color: #111; font-size: 10pt; }}
.report-head {{ display:flex; justify-content:space-between; gap:16px; border-bottom:2px solid #111; padding-bottom:10px; margin-bottom:14px; }}
.logo {{ width:90px; height:50px; border:1px dashed #777; display:flex; align-items:center; justify-content:center; font-size:8pt; }}
.grid {{ display:grid; grid-template-columns:repeat(12, 1fr); gap:10px; }}
.grid section {{ border:1px solid #bbb; padding:8px; break-inside:avoid; }}
{''.join(f'.span-{i} {{ grid-column: span {i}; }}' for i in range(1, 13))}
h1 {{ margin:0; font-size:18pt; }} h2 {{ margin:0 0 6px; font-size:11pt; }}
table {{ width:100%; border-collapse:collapse; }} th,td {{ border:1px solid #ccc; padding:4px; vertical-align:top; }} th {{ background:#f3f3f3; }}
.page-break {{ break-before:page; grid-column:1/-1; }} .signatures {{ display:grid; grid-template-columns:repeat(3,1fr); gap:20px; padding-top:28px; }}
.meta {{ font-size:8pt; color:#555; }} .empty {{ color:#777; font-style:italic; }}
</style></head><body>
<div class="meta">{escape(header_html)}</div>
<div class="report-head"><div>{logo_marker}<strong>{company}</strong></div><div><h1>{title}</h1><div>{escape(str(project.get('number') or ''))} · {escape(str(project.get('name') or ''))}</div><div>{escape(str(report.get('report_date') or ''))} · {escape(str(report.get('shift_code') or ''))}</div></div></div>
<div class="grid">{''.join(blocks)}</div>
<div class="meta">{escape(footer_html)}</div>
</body></html>"""
    return html.encode("utf-8")


def _json(snapshot: Mapping[str, Any], layout: DPRLayoutSpec) -> bytes:
    payload = {"data": snapshot, "presentation": layout.model_dump(mode="json")}
    return json.dumps(payload, indent=2, ensure_ascii=False, sort_keys=True).encode("utf-8")


def _csv(snapshot: Mapping[str, Any], layout: DPRLayoutSpec) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["section", "row", "field", "value"])
    for element in _ordered_elements(layout):
        key = element.kind.value
        if element.kind in {DPRLayoutElementKind.PAGE_BREAK, DPRLayoutElementKind.CUSTOM_TEXT, DPRLayoutElementKind.SIGNATURES}:
            continue
        if element.kind == DPRLayoutElementKind.PROJECT_SUMMARY:
            rows = [{"field": name, "value": value} for name, value in _project_summary(snapshot)]
            fields = ["field", "value"]
        elif element.kind == DPRLayoutElementKind.WEATHER:
            rows = [{"field": name, "value": value} for name, value in _weather(snapshot)]
            fields = ["field", "value"]
        elif element.kind == DPRLayoutElementKind.NOTES:
            report = snapshot.get("report", {}) if isinstance(snapshot.get("report"), Mapping) else {}
            rows = [{"notes": report.get("notes")}]
            fields = ["notes"]
        else:
            rows = _rows(snapshot, key)
            fields = _field_list(element, key, rows)
        for index, row in enumerate(rows, start=1):
            for field in fields:
                writer.writerow([key, index, field, _scalar(row.get(field))])
    return output.getvalue().encode("utf-8")


def _xlsx(snapshot: Mapping[str, Any], layout: DPRLayoutSpec) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()
    for element in _ordered_elements(layout):
        if element.kind in {DPRLayoutElementKind.PAGE_BREAK, DPRLayoutElementKind.CUSTOM_TEXT, DPRLayoutElementKind.SIGNATURES}:
            continue
        base_name = (element.title or element.kind.value).strip()[:25] or "Section"
        name = base_name
        suffix = 2
        while name in used_names:
            name = f"{base_name[:22]} {suffix}"
            suffix += 1
        used_names.add(name)
        sheet = workbook.create_sheet(name)
        key = element.kind.value
        if element.kind == DPRLayoutElementKind.PROJECT_SUMMARY:
            rows = [{"field": field, "value": value} for field, value in _project_summary(snapshot)]
            fields = ["field", "value"]
        elif element.kind == DPRLayoutElementKind.WEATHER:
            rows = [{"field": field, "value": value} for field, value in _weather(snapshot)]
            fields = ["field", "value"]
        elif element.kind == DPRLayoutElementKind.NOTES:
            report = snapshot.get("report", {}) if isinstance(snapshot.get("report"), Mapping) else {}
            rows = [{"notes": report.get("notes")}]
            fields = ["notes"]
        else:
            rows = _rows(snapshot, key)
            fields = _field_list(element, key, rows)
        for col, field in enumerate(fields, start=1):
            sheet.cell(row=1, column=col, value=_label(field))
        for row_index, row in enumerate(rows, start=2):
            for col, field in enumerate(fields, start=1):
                sheet.cell(row=row_index, column=col, value=_scalar(row.get(field)))
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = min(60, max(12, max((len(str(cell.value or "")) for cell in column), default=12) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    if not workbook.worksheets:
        workbook.create_sheet("DPR")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx(snapshot: Mapping[str, Any], layout: DPRLayoutSpec) -> bytes:
    document = Document()
    document.add_heading(_company_name(snapshot, layout), level=2)
    document.add_heading(layout.branding.report_title, level=1)
    for element in _ordered_elements(layout):
        if element.kind == DPRLayoutElementKind.PAGE_BREAK:
            document.add_page_break()
            continue
        if element.start_new_page:
            document.add_page_break()
        document.add_heading(element.title or _label(element.kind.value), level=2)
        if element.kind == DPRLayoutElementKind.CUSTOM_TEXT:
            document.add_paragraph(element.custom_text or "")
            continue
        if element.kind == DPRLayoutElementKind.SIGNATURES:
            table = document.add_table(rows=2, cols=3)
            for index, text in enumerate(("Prepared by", "Checked by", "Approved by")):
                table.cell(0, index).text = text
                table.cell(1, index).text = "\n\n________________________"
            continue
        key = element.kind.value
        if element.kind == DPRLayoutElementKind.PROJECT_SUMMARY:
            rows = [{"field": name, "value": value} for name, value in _project_summary(snapshot)]
            fields = ["field", "value"]
        elif element.kind == DPRLayoutElementKind.WEATHER:
            rows = [{"field": name, "value": value} for name, value in _weather(snapshot)]
            fields = ["field", "value"]
        elif element.kind == DPRLayoutElementKind.NOTES:
            report = snapshot.get("report", {}) if isinstance(snapshot.get("report"), Mapping) else {}
            document.add_paragraph(_scalar(report.get("notes")))
            continue
        else:
            rows = _rows(snapshot, key)
            fields = _field_list(element, key, rows)
        if not fields:
            document.add_paragraph("No entries")
            continue
        table = document.add_table(rows=1, cols=len(fields))
        for index, field in enumerate(fields):
            table.rows[0].cells[index].text = _label(field)
        for row in rows:
            cells = table.add_row().cells
            for index, field in enumerate(fields):
                cells[index].text = _scalar(row.get(field))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf(snapshot: Mapping[str, Any], layout: DPRLayoutSpec) -> bytes:
    page_sizes = {"a4": A4, "a3": A3, "letter": LETTER}
    size = page_sizes[layout.page.size.value]
    if layout.page.orientation.value == "landscape":
        size = landscape(size)
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=size,
        topMargin=layout.page.margin_top_mm * mm,
        rightMargin=layout.page.margin_right_mm * mm,
        bottomMargin=layout.page.margin_bottom_mm * mm,
        leftMargin=layout.page.margin_left_mm * mm,
    )
    styles = getSampleStyleSheet()
    story: list[Any] = [
        Paragraph(escape(_company_name(snapshot, layout)), styles["Heading2"]),
        Paragraph(escape(layout.branding.report_title), styles["Title"]),
        Spacer(1, 6),
    ]
    for element in _ordered_elements(layout):
        if element.kind == DPRLayoutElementKind.PAGE_BREAK:
            story.append(PageBreak())
            continue
        if element.start_new_page:
            story.append(PageBreak())
        story.append(Paragraph(escape(element.title or _label(element.kind.value)), styles["Heading3"]))
        if element.kind == DPRLayoutElementKind.CUSTOM_TEXT:
            story.append(Paragraph(escape(element.custom_text or ""), styles["BodyText"]))
            continue
        if element.kind == DPRLayoutElementKind.SIGNATURES:
            table = Table([["Prepared by", "Checked by", "Approved by"], ["\n\n____________", "\n\n____________", "\n\n____________"]])
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)]))
            story.append(table)
            continue
        key = element.kind.value
        if element.kind == DPRLayoutElementKind.PROJECT_SUMMARY:
            rows = [{"field": name, "value": value} for name, value in _project_summary(snapshot)]
            fields = ["field", "value"]
        elif element.kind == DPRLayoutElementKind.WEATHER:
            rows = [{"field": name, "value": value} for name, value in _weather(snapshot)]
            fields = ["field", "value"]
        elif element.kind == DPRLayoutElementKind.NOTES:
            report = snapshot.get("report", {}) if isinstance(snapshot.get("report"), Mapping) else {}
            story.append(Paragraph(escape(_scalar(report.get("notes"))), styles["BodyText"]))
            continue
        else:
            rows = _rows(snapshot, key)
            fields = _field_list(element, key, rows)
        if not fields:
            story.append(Paragraph("No entries", styles["BodyText"]))
            continue
        table_data = [[_label(field) for field in fields]]
        table_data.extend([[_scalar(row.get(field)) for field in fields] for row in rows])
        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 8))
    document.build(story)
    return buffer.getvalue()


def render_dpr(
    snapshot: Mapping[str, Any],
    layout: DPRLayoutSpec,
    output_format: DPROutputFormat,
) -> tuple[bytes, str, str]:
    if output_format == DPROutputFormat.HTML:
        content = _html(snapshot, layout)
    elif output_format == DPROutputFormat.JSON:
        content = _json(snapshot, layout)
    elif output_format == DPROutputFormat.CSV:
        content = _csv(snapshot, layout)
    elif output_format == DPROutputFormat.XLSX:
        content = _xlsx(snapshot, layout)
    elif output_format == DPROutputFormat.DOCX:
        content = _docx(snapshot, layout)
    elif output_format == DPROutputFormat.PDF:
        content = _pdf(snapshot, layout)
    else:
        raise ValueError(f"Unsupported DPR output format: {output_format}")
    return content, _CONTENT_TYPES[output_format], _EXTENSIONS[output_format]
