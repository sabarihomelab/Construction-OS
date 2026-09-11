import asyncio
import io
import re
import shutil
import subprocess
import tempfile
from collections.abc import Mapping
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.files.generated import read_file_version_bytes
from app.modules.reporting.template_models import (
    ReportTemplateVersion,
    TemplateOutputFormat,
    TemplateSourceKind,
)
from app.modules.reporting.template_provider import resolve_path
from app.modules.reporting.template_renderer import render_report, visible_blocks
from app.modules.reporting.template_schemas import ReportBlockKind, ReportLayoutSpec

_PLACEHOLDER = re.compile(r"{{\s*([A-Za-z0-9_.-]+)\s*}}")


class ReportExecutionError(ValueError):
    pass


def _string(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def _replace_text(text: str, payload: Mapping[str, object]) -> str:
    return _PLACEHOLDER.sub(lambda match: _string(resolve_path(payload, match.group(1))), text)


def _replace_paragraph(paragraph: Paragraph, payload: Mapping[str, object]) -> None:
    original = paragraph.text
    replaced = _replace_text(original, payload)
    if original == replaced:
        return
    if paragraph.runs:
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replaced)


def _replace_table_scalars(table: Table, payload: Mapping[str, object]) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                _replace_paragraph(paragraph, payload)
            for nested in cell.tables:
                _replace_table_scalars(nested, payload)


def _collection_rows(payload: Mapping[str, object], key: str) -> list[dict[str, object]]:
    value = resolve_path(payload, key)
    if not isinstance(value, list):
        return []
    return [dict(item) for item in value if isinstance(item, Mapping)]


def _insert_collection_table(
    document: Any,
    marker: Paragraph,
    *,
    rows: list[dict[str, object]],
    columns: list[tuple[str, str]],
) -> None:
    if not rows or not columns:
        marker.text = ""
        return
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for index, (_, label) in enumerate(columns):
        table.rows[0].cells[index].text = label
    for item in rows:
        cells = table.add_row().cells
        for index, (key, _) in enumerate(columns):
            cells[index].text = _string(item.get(key))
    marker._p.addnext(table._tbl)  # noqa: SLF001
    marker._element.getparent().remove(marker._element)  # noqa: SLF001


def render_uploaded_docx(
    source: bytes,
    payload: Mapping[str, object],
    layout: ReportLayoutSpec,
) -> bytes:
    try:
        document = Document(io.BytesIO(source))
    except Exception as exc:
        raise ReportExecutionError("Uploaded Word template is not a readable DOCX file") from exc

    visible_tables = {
        block.data_key: block
        for block in visible_blocks(payload, layout)
        if block.kind == ReportBlockKind.TABLE and block.data_key
    }
    body_markers = list(document.paragraphs)
    for marker in body_markers:
        match = _PLACEHOLDER.fullmatch(marker.text.strip())
        if match and match.group(1) in {item.data_key for item in layout.blocks if item.data_key}:
            key = match.group(1)
            block = visible_tables.get(key)
            if block is None:
                marker.text = ""
                continue
            rows = _collection_rows(payload, key)
            configured = sorted(
                (column for column in block.table_columns if column.visible),
                key=lambda column: (column.order, column.key),
            )
            if configured:
                columns = [(column.key, column.label or column.key.replace("_", " ").title()) for column in configured]
            elif rows:
                columns = [
                    (field, field.replace("_", " ").title())
                    for field in rows[0]
                    if field not in {"id", "source", "source_id", "source_revision"}
                ]
            else:
                columns = []
            _insert_collection_table(document, marker, rows=rows, columns=columns)
        else:
            _replace_paragraph(marker, payload)

    for table in document.tables:
        _replace_table_scalars(table, payload)
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            _replace_paragraph(paragraph, payload)
        for table in section.header.tables:
            _replace_table_scalars(table, payload)
        for paragraph in section.footer.paragraphs:
            _replace_paragraph(paragraph, payload)
        for table in section.footer.tables:
            _replace_table_scalars(table, payload)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _convert_docx_to_pdf_sync(content: bytes) -> bytes:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        raise ReportExecutionError(
            "Exact uploaded Word-template PDF conversion requires LibreOffice/soffice on the document worker"
        )
    with tempfile.TemporaryDirectory(prefix="construction-os-report-") as directory:
        root = Path(directory)
        source = root / "report.docx"
        source.write_bytes(content)
        completed = subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(root),
                str(source),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=120,
        )
        pdf = root / "report.pdf"
        if completed.returncode != 0 or not pdf.is_file():
            message = (completed.stderr or completed.stdout or "Word to PDF conversion failed").strip()
            raise ReportExecutionError(message[:1000])
        return pdf.read_bytes()


async def render_template_version(
    db: AsyncSession,
    *,
    organization_id,
    template_version: ReportTemplateVersion,
    payload: dict[str, object],
    layout: ReportLayoutSpec,
    output_format: TemplateOutputFormat,
) -> tuple[bytes, str, str]:
    is_uploaded_docx = (
        template_version.source_kind == TemplateSourceKind.UPLOADED
        and (template_version.source_format or "").lower() == "docx"
        and template_version.source_file_asset_id is not None
        and template_version.source_file_version is not None
    )
    if not is_uploaded_docx:
        return render_report(payload, layout, output_format)

    _version, source = await read_file_version_bytes(
        db,
        organization_id=organization_id,
        asset_id=template_version.source_file_asset_id,
        version=template_version.source_file_version,
    )
    filled_docx = render_uploaded_docx(source, payload, layout)
    if output_format == TemplateOutputFormat.DOCX:
        return (
            filled_docx,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )
    if output_format == TemplateOutputFormat.PDF:
        pdf = await asyncio.to_thread(_convert_docx_to_pdf_sync, filled_docx)
        return pdf, "application/pdf", "pdf"
    raise ReportExecutionError(
        "Uploaded Word templates support DOCX and PDF output; use a designer template for other formats"
    )
