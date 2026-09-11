import csv
import io
import json
from collections.abc import Callable, Mapping, Sequence
from html import escape
from typing import Any

from docx import Document
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, LETTER, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    KeepTogether,
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.modules.reporting.template_models import TemplateOutputFormat
from app.modules.reporting.template_provider import resolve_path
from app.modules.reporting.template_schemas import (
    ConditionOperator,
    ReportBlock,
    ReportBlockKind,
    ReportLayoutSpec,
)

ImageLoader = Callable[[object], bytes | None]

_CONTENT_TYPES = {
    TemplateOutputFormat.PDF: "application/pdf",
    TemplateOutputFormat.HTML: "text/html; charset=utf-8",
    TemplateOutputFormat.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    TemplateOutputFormat.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    TemplateOutputFormat.CSV: "text/csv; charset=utf-8",
    TemplateOutputFormat.JSON: "application/json; charset=utf-8",
}


def _scalar(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _label(value: str) -> str:
    return value.replace("_", " ").replace(".", " / ").strip().title()


def _truthy_non_empty(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, Mapping | Sequence) and not isinstance(value, (str, bytes, bytearray)):
        return len(value) > 0
    return True


def _condition_matches(payload: Mapping[str, object], block: ReportBlock) -> bool:
    condition = block.condition
    if condition is None:
        return True
    actual = resolve_path(payload, condition.path)
    operator = condition.operator
    expected = condition.value
    if operator == ConditionOperator.EXISTS:
        return actual is not None
    if operator == ConditionOperator.NON_EMPTY:
        return _truthy_non_empty(actual)
    if operator == ConditionOperator.EQUALS:
        return actual == expected
    if operator == ConditionOperator.NOT_EQUALS:
        return actual != expected
    if operator == ConditionOperator.CONTAINS:
        if isinstance(actual, str):
            return str(expected) in actual
        if isinstance(actual, Sequence) and not isinstance(actual, (str, bytes, bytearray)):
            return expected in actual
        return False
    try:
        if operator == ConditionOperator.GREATER_THAN:
            return bool(actual > expected)  # type: ignore[operator]
        if operator == ConditionOperator.GREATER_OR_EQUAL:
            return bool(actual >= expected)  # type: ignore[operator]
        if operator == ConditionOperator.LESS_THAN:
            return bool(actual < expected)  # type: ignore[operator]
        if operator == ConditionOperator.LESS_OR_EQUAL:
            return bool(actual <= expected)  # type: ignore[operator]
    except TypeError:
        return False
    return False


def _resolve_text(template: str | None, payload: Mapping[str, object]) -> str:
    if not template:
        return ""
    result = template
    from app.modules.reporting.template_provider import placeholders_in_text

    for path in placeholders_in_text(template):
        result = result.replace("{{" + path + "}}", _scalar(resolve_path(payload, path)))
        result = result.replace("{{ " + path + " }}", _scalar(resolve_path(payload, path)))
    return result


def _collection(payload: Mapping[str, object], key: str | None) -> list[dict[str, object]]:
    if not key:
        return []
    value = resolve_path(payload, key)
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes, bytearray)):
        return []
    return [dict(item) for item in value if isinstance(item, Mapping)]


def _sort_rows(rows: list[dict[str, object]], block: ReportBlock) -> list[dict[str, object]]:
    result = list(rows)
    for rule in reversed(block.sort_by):
        result.sort(
            key=lambda row: (row.get(rule.key) is None, str(row.get(rule.key) or "")),
            reverse=rule.direction == "desc",
        )
    return result


def _columns(block: ReportBlock, rows: list[dict[str, object]]) -> list[tuple[str, str]]:
    configured = [item for item in block.table_columns if item.visible]
    if configured:
        configured.sort(key=lambda item: (item.order, item.key))
        return [(item.key, item.label or _label(item.key)) for item in configured]
    if not rows:
        return []
    return [(key, _label(key)) for key in rows[0].keys() if key not in {"id", "source"}]


def _block_empty(payload: Mapping[str, object], block: ReportBlock) -> bool:
    if block.kind == ReportBlockKind.TABLE:
        return not _collection(payload, block.data_key)
    if block.kind == ReportBlockKind.FIELD_GROUP:
        return not any(_truthy_non_empty(resolve_path(payload, path)) for path in block.scalar_fields)
    if block.kind == ReportBlockKind.TEXT:
        return not bool(_resolve_text(block.text_template, payload).strip())
    if block.kind == ReportBlockKind.IMAGE:
        return resolve_path(payload, block.image_placeholder or "") is None
    return False


def visible_blocks(payload: Mapping[str, object], layout: ReportLayoutSpec) -> list[ReportBlock]:
    blocks = sorted(layout.blocks, key=lambda item: (item.row, item.order, item.id))
    return [
        block
        for block in blocks
        if block.visible
        and _condition_matches(payload, block)
        and not (block.hide_when_empty and _block_empty(payload, block))
    ]


def _html(payload: Mapping[str, object], layout: ReportLayoutSpec) -> bytes:
    blocks_html: list[str] = []
    for block in visible_blocks(payload, layout):
        if block.kind == ReportBlockKind.PAGE_BREAK:
            blocks_html.append('<div class="page-break"></div>')
            continue
        if block.page_break_before:
            blocks_html.append('<div class="page-break"></div>')
        title = escape(_resolve_text(block.title, payload)) if block.title else ""
        title_html = f"<h2>{title}</h2>" if title else ""
        content = ""
        if block.kind == ReportBlockKind.FIELD_GROUP:
            cells = []
            for path in block.scalar_fields:
                cells.append(
                    f'<div class="field"><span>{escape(_label(path.split(".")[-1]))}</span>'
                    f'<strong>{escape(_scalar(resolve_path(payload, path)))}</strong></div>'
                )
            content = f'<div class="field-grid">{"".join(cells)}</div>'
        elif block.kind == ReportBlockKind.TABLE:
            rows = _sort_rows(_collection(payload, block.data_key), block)
            columns = _columns(block, rows)
            headers = "".join(f"<th>{escape(label)}</th>" for _, label in columns)
            body = []
            for row in rows:
                cells = "".join(f"<td>{escape(_scalar(row.get(key)))}</td>" for key, _ in columns)
                body.append(f"<tr>{cells}</tr>")
            content = (
                f'<table><thead><tr>{headers}</tr></thead><tbody>{"".join(body)}</tbody></table>'
                if columns
                else '<p class="empty">No entries</p>'
            )
        elif block.kind == ReportBlockKind.TEXT:
            content = f"<p>{escape(_resolve_text(block.text_template, payload))}</p>"
        elif block.kind == ReportBlockKind.IMAGE:
            value = resolve_path(payload, block.image_placeholder or "")
            src = ""
            if isinstance(value, str) and value.startswith(("data:image/", "http://", "https://", "/")):
                src = value
            elif isinstance(value, Mapping):
                candidate = value.get("url") or value.get("data_uri")
                if isinstance(candidate, str):
                    src = candidate
            if src:
                content = f'<img class="report-image" src="{escape(src)}" alt="{escape(title)}">'
        elif block.kind == ReportBlockKind.SIGNATURES:
            signatures = resolve_path(payload, block.data_key or "approvals")
            if isinstance(signatures, Mapping):
                cells = [
                    f'<div><span>{escape(_label(str(key)))}</span><strong>{escape(_scalar(value))}</strong></div>'
                    for key, value in signatures.items()
                ]
                content = f'<div class="signature-grid">{"".join(cells)}</div>'
        blocks_html.append(
            f'<section style="width:{block.width_percent}%;" class="report-block">{title_html}{content}</section>'
        )

    page = layout.page
    title = escape(_resolve_text(layout.report_title, payload))
    hf = layout.header_footer
    header = " | ".join(
        item
        for item in (
            _resolve_text(hf.header_left, payload),
            _resolve_text(hf.header_center, payload),
            _resolve_text(hf.header_right, payload),
        )
        if item
    )
    footer = " | ".join(
        item
        for item in (
            _resolve_text(hf.footer_left, payload),
            _resolve_text(hf.footer_center, payload),
            _resolve_text(hf.footer_right, payload),
        )
        if item
    )
    css_page = f"{page.size.value.upper()} {page.orientation.value}"
    html = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{title}</title><style>
@page {{ size:{css_page}; margin:{page.margin_top_mm}mm {page.margin_right_mm}mm {page.margin_bottom_mm}mm {page.margin_left_mm}mm; }}
body {{ font-family:Arial,sans-serif; color:#111; font-size:10pt; }}
.report-title {{ border-bottom:2px solid #111; padding-bottom:8px; margin-bottom:12px; }}
.report-grid {{ display:flex; flex-wrap:wrap; align-items:flex-start; }}
.report-block {{ box-sizing:border-box; padding:6px; break-inside:avoid; }}
.report-block h2 {{ font-size:11pt; margin:0 0 6px; break-after:avoid; }}
.field-grid {{ display:grid; grid-template-columns:repeat(2,minmax(0,1fr)); gap:5px 14px; }}
.field span {{ display:block; font-size:8pt; color:#666; }}
table {{ width:100%; border-collapse:collapse; }}
thead {{ display:table-header-group; }} tr {{ break-inside:avoid; }}
th,td {{ border:1px solid #bbb; padding:4px; vertical-align:top; }} th {{ background:#f0f0f0; }}
.page-break {{ width:100%; break-before:page; }}
.signature-grid {{ display:grid; grid-template-columns:repeat(3,minmax(0,1fr)); gap:18px; padding-top:24px; }}
.signature-grid span {{ display:block; color:#666; font-size:8pt; }}
.report-image {{ max-width:100%; max-height:120px; object-fit:contain; }}
.meta {{ font-size:8pt; color:#555; }} .empty {{ color:#777; font-style:italic; }}
@media print {{ .no-print {{ display:none !important; }} }}
</style></head><body>
<div class="meta">{escape(header)}</div><h1 class="report-title">{title}</h1>
<div class="report-grid">{"".join(blocks_html)}</div><div class="meta">{escape(footer)}</div>
</body></html>"""
    return html.encode("utf-8")


def _pdf(
    payload: Mapping[str, object],
    layout: ReportLayoutSpec,
    image_loader: ImageLoader | None,
) -> bytes:
    page_sizes = {"a4": A4, "a3": A3, "letter": LETTER}
    size = page_sizes[layout.page.size.value]
    if layout.page.orientation.value == "landscape":
        size = landscape(size)
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=size,
        topMargin=layout.page.margin_top_mm * mm,
        rightMargin=layout.page.margin_right_mm * mm,
        bottomMargin=layout.page.margin_bottom_mm * mm,
        leftMargin=layout.page.margin_left_mm * mm,
    )
    styles = getSampleStyleSheet()
    section_heading = ParagraphStyle(
        "ReportSection",
        parent=styles["Heading3"],
        keepWithNext=True,
        spaceBefore=6,
        spaceAfter=4,
    )
    story: list[Any] = [Paragraph(escape(_resolve_text(layout.report_title, payload)), styles["Title"])]
    for block in visible_blocks(payload, layout):
        if block.kind == ReportBlockKind.PAGE_BREAK:
            story.append(PageBreak())
            continue
        if block.page_break_before:
            story.append(PageBreak())
        pieces: list[Any] = []
        if block.title:
            pieces.append(Paragraph(escape(_resolve_text(block.title, payload)), section_heading))
        if block.kind == ReportBlockKind.FIELD_GROUP:
            rows = [
                [
                    Paragraph(escape(_label(path.split(".")[-1])), styles["BodyText"]),
                    Paragraph(escape(_scalar(resolve_path(payload, path))), styles["BodyText"]),
                ]
                for path in block.scalar_fields
            ]
            table = Table(rows, colWidths=[45 * mm, None]) if rows else None
            if table:
                table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
                pieces.append(table)
        elif block.kind == ReportBlockKind.TABLE:
            rows = _sort_rows(_collection(payload, block.data_key), block)
            columns = _columns(block, rows)
            if columns:
                table_data = [[Paragraph(escape(label), styles["BodyText"]) for _, label in columns]]
                for row in rows:
                    table_data.append(
                        [Paragraph(escape(_scalar(row.get(key))), styles["BodyText"]) for key, _ in columns]
                    )
                table = LongTable(
                    table_data,
                    repeatRows=1 if block.repeat_table_header else 0,
                    splitByRow=1 if block.avoid_row_split else 0,
                )
                table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("FONTSIZE", (0, 0), (-1, -1), 7),
                            ("LEFTPADDING", (0, 0), (-1, -1), 3),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                        ]
                    )
                )
                pieces.append(table)
        elif block.kind == ReportBlockKind.TEXT:
            pieces.append(Paragraph(escape(_resolve_text(block.text_template, payload)), styles["BodyText"]))
        elif block.kind == ReportBlockKind.IMAGE and image_loader:
            value = resolve_path(payload, block.image_placeholder or "")
            binary = image_loader(value)
            if binary:
                pieces.append(Image(io.BytesIO(binary), width=45 * mm, height=25 * mm, kind="proportional"))
        elif block.kind == ReportBlockKind.SIGNATURES:
            signatures = resolve_path(payload, block.data_key or "approvals")
            if isinstance(signatures, Mapping):
                keys = list(signatures.keys())
                table = Table(
                    [[_label(str(key)) for key in keys], [_scalar(signatures[key]) for key in keys]],
                    repeatRows=1,
                )
                table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey)]))
                pieces.append(table)
        if block.keep_title_with_content and len(pieces) <= 2:
            story.append(KeepTogether(pieces))
        else:
            story.extend(pieces)
        story.append(Spacer(1, 6))
    document.build(story)
    return buffer.getvalue()


def _docx(payload: Mapping[str, object], layout: ReportLayoutSpec) -> bytes:
    document = Document()
    document.add_heading(_resolve_text(layout.report_title, payload), level=1)
    for block in visible_blocks(payload, layout):
        if block.kind == ReportBlockKind.PAGE_BREAK or block.page_break_before:
            document.add_page_break()
            if block.kind == ReportBlockKind.PAGE_BREAK:
                continue
        if block.title:
            document.add_heading(_resolve_text(block.title, payload), level=2)
        if block.kind == ReportBlockKind.FIELD_GROUP:
            table = document.add_table(rows=0, cols=2)
            for path in block.scalar_fields:
                cells = table.add_row().cells
                cells[0].text = _label(path.split(".")[-1])
                cells[1].text = _scalar(resolve_path(payload, path))
        elif block.kind == ReportBlockKind.TABLE:
            rows = _sort_rows(_collection(payload, block.data_key), block)
            columns = _columns(block, rows)
            if columns:
                table = document.add_table(rows=1, cols=len(columns))
                for index, (_, label) in enumerate(columns):
                    table.rows[0].cells[index].text = label
                for row in rows:
                    cells = table.add_row().cells
                    for index, (key, _) in enumerate(columns):
                        cells[index].text = _scalar(row.get(key))
        elif block.kind == ReportBlockKind.TEXT:
            document.add_paragraph(_resolve_text(block.text_template, payload))
        elif block.kind == ReportBlockKind.SIGNATURES:
            signatures = resolve_path(payload, block.data_key or "approvals")
            if isinstance(signatures, Mapping):
                table = document.add_table(rows=2, cols=len(signatures))
                for index, (key, value) in enumerate(signatures.items()):
                    table.cell(0, index).text = _label(str(key))
                    table.cell(1, index).text = _scalar(value)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx(payload: Mapping[str, object], layout: ReportLayoutSpec) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for block in visible_blocks(payload, layout):
        if block.kind not in {ReportBlockKind.TABLE, ReportBlockKind.FIELD_GROUP}:
            continue
        base = (_resolve_text(block.title, payload) or block.id or "Section")[:25]
        name = base
        counter = 2
        while name in used:
            name = f"{base[:22]} {counter}"
            counter += 1
        used.add(name)
        sheet = workbook.create_sheet(name)
        if block.kind == ReportBlockKind.FIELD_GROUP:
            for index, path in enumerate(block.scalar_fields, start=1):
                sheet.cell(index, 1, _label(path.split(".")[-1]))
                sheet.cell(index, 2, _scalar(resolve_path(payload, path)))
        else:
            rows = _sort_rows(_collection(payload, block.data_key), block)
            columns = _columns(block, rows)
            for column_index, (_, label) in enumerate(columns, start=1):
                sheet.cell(1, column_index, label)
            for row_index, row in enumerate(rows, start=2):
                for column_index, (key, _) in enumerate(columns, start=1):
                    sheet.cell(row_index, column_index, _scalar(row.get(key)))
            sheet.freeze_panes = "A2"
    if not workbook.worksheets:
        workbook.create_sheet("Report")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _csv(payload: Mapping[str, object], layout: ReportLayoutSpec) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["block", "row", "field", "value"])
    for block in visible_blocks(payload, layout):
        if block.kind == ReportBlockKind.FIELD_GROUP:
            for path in block.scalar_fields:
                writer.writerow([block.id, 1, path, _scalar(resolve_path(payload, path))])
        elif block.kind == ReportBlockKind.TABLE:
            rows = _sort_rows(_collection(payload, block.data_key), block)
            columns = _columns(block, rows)
            for row_index, row in enumerate(rows, start=1):
                for key, label in columns:
                    writer.writerow([block.id, row_index, label, _scalar(row.get(key))])
    return output.getvalue().encode("utf-8")


def _json(payload: Mapping[str, object], layout: ReportLayoutSpec) -> bytes:
    return json.dumps(
        {"payload": payload, "presentation": layout.model_dump(mode="json")},
        ensure_ascii=False,
        indent=2,
        sort_keys=True,
    ).encode("utf-8")


def render_report(
    payload: Mapping[str, object],
    layout: ReportLayoutSpec,
    output_format: TemplateOutputFormat,
    *,
    image_loader: ImageLoader | None = None,
) -> tuple[bytes, str, str]:
    if output_format == TemplateOutputFormat.HTML:
        content = _html(payload, layout)
    elif output_format == TemplateOutputFormat.PDF:
        content = _pdf(payload, layout, image_loader)
    elif output_format == TemplateOutputFormat.XLSX:
        content = _xlsx(payload, layout)
    elif output_format == TemplateOutputFormat.DOCX:
        content = _docx(payload, layout)
    elif output_format == TemplateOutputFormat.CSV:
        content = _csv(payload, layout)
    elif output_format == TemplateOutputFormat.JSON:
        content = _json(payload, layout)
    else:
        raise ValueError(f"Unsupported report output format: {output_format}")
    return content, _CONTENT_TYPES[output_format], output_format.value
